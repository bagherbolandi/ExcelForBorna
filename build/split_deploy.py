#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Borna PM — Split deployment generator (دسترسی واقعی روی شبکه داخلی)
--------------------------------------------------------------------
از فایل مادر (out/Project_Management_Borna.xlsx + .meta.json) برای هر واحد
یک فایل فرم مستقل می‌سازد:

  * شیت‌های ورودی واحد = فرمول‌های زنده (همان فرمول‌های فایل مادر؛ ارجاع به
    شیت‌های بالادستی به‌جای داده واقعی، اسنپ‌شات مقادیر با نامِ همان شیت است).
  * شیت‌های سایر واحدها = اسنپ‌شاتِ فقط‌خواندنی و مخفی (مقادیر محاسبه‌شده)؛
    جدول کاربران و توکن‌ها در این فایل‌ها پاک می‌شوند.
  * کل شیت‌ها با «توکن واحد» قفل می‌شوند؛ ساختار فایل قفل است؛
    واحد فقط سلول‌های unlocked شیت خودش را می‌تواند ویرایش کند — و
    با ACL شبکه حتی فایل واحدهای دیگر را هم باز نمی‌کند.

چرخه کار (ماکروهای master):  Borna_Collect → بازبینی PM → Borna_Publish

Usage:
    python3 build/split_deploy.py --master out/Project_Management_Borna.xlsx --out out/split
"""
import argparse
import datetime as dt
import json
import os
import re
from copy import copy, deepcopy

from openpyxl import Workbook, load_workbook
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.workbook.protection import WorkbookProtection
from openpyxl.worksheet.datavalidation import DataValidation
import formulas

ROLE_OF_DEPT = {"INTAKE": ["INTAKE"], "ENG": ["ENG"], "TRD": ["TRD_IN", "TRD_EX"],
                "PLAN": ["PLAN"], "FIN": ["FIN"], "SALES": ["SALES"], "EXEC": ["CEO"]}
ALWAYS_VISIBLE = {"راهنما"}
SKIP_SHEETS = {"ورود"}


def clean(v):
    if v is None:
        return ""
    if isinstance(v, (int, float, str, bool, dt.date, dt.datetime)):
        return v
    try:
        import numpy as _np
        if isinstance(v, _np.generic):
            return v.item()
    except Exception:
        pass
    try:
        if hasattr(v, "value"):
            v = v.value[0, 0]
            return clean(v)
    except Exception:
        pass
    return str(v)


def copy_style(src, dst, lock_all=False):
    dst.font = copy(src.font)
    if src.has_style:
        dst.fill = copy(src.fill)
        dst.border = copy(src.border)
        dst.alignment = copy(src.alignment)
        dst.number_format = src.number_format
    if lock_all:
        from openpyxl.styles import Protection
        dst.protection = Protection(locked=True)
    else:
        dst.protection = copy(src.protection)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--master", default="out/Project_Management_Borna.xlsx")
    ap.add_argument("--out", default="out/split")
    args = ap.parse_args()

    meta = json.load(open(args.master + ".meta.json", encoding="utf-8"))
    ACCESS = meta["access"]; SHEET_ORDER = meta["sheet_order"]
    U0, U1 = meta["settings"]["u_rows"]
    CH = meta["settings"]["conn"]["hdr"]; CF_, CL_ = meta["settings"]["conn"]["first"], meta["settings"]["conn"]["last"]
    TOKENS = meta["dept_tokens"]
    DEPTS = {d[0]: d for d in meta["depts"]}

    mwb = load_workbook(args.master)                      # formulas + styles
    fname = os.path.basename(args.master)
    print("evaluating master with formulas engine …")
    sol = formulas.ExcelModel().loads(args.master).finish().calculate()
    def engine_val(sheet, coord):
        return sol.get(f"'[{fname}]{sheet}'!{coord}")

    os.makedirs(args.out, exist_ok=True)
    summary = []
    for dkey, (dk, folder, dfile, _cols) in [(k, DEPTS[k]) for k in DEPTS]:
        roles = ROLE_OF_DEPT[dk]
        input_sheets = [s for s in SHEET_ORDER if s not in SKIP_SHEETS
                        and any(ACCESS.get(s, {}).get(r) == "W" for r in roles)]
        snap_sheets = [s for s in SHEET_ORDER if s not in SKIP_SHEETS and s not in input_sheets]

        dwb = Workbook()
        dwb.calculation.fullCalcOnLoad = True
        dwb.remove(dwb.active)
        first = True
        for s in SHEET_ORDER:
            if s in SKIP_SHEETS:
                continue
            src = mwb[s]
            is_input = s in input_sheets
            dst = dwb.create_sheet(s)
            if first:
                first = False
            maxr, maxc = src.max_row, min(src.max_column, 30)
            # cells
            for row in src.iter_rows(min_row=1, max_row=maxr, min_col=1, max_col=maxc):
                for sc in row:
                    dc = dst.cell(row=sc.row, column=sc.column)
                    v = sc.value
                    if is_input:
                        dc.value = v
                    else:
                        if s == "تنظیمات" and (
                                (U0 - 1 <= sc.row <= U1 and sc.column <= 6) or
                                (CH - 1 <= sc.row <= CL_ and sc.column <= 6)):
                            v = None  # mask users / dept tokens
                        if isinstance(v, str) and v.startswith("="):
                            dc.value = clean(engine_val(s, sc.coordinate))
                        else:
                            dc.value = v
                    copy_style(sc, dc, lock_all=not is_input)
            # merges, widths, heights
            for rng in list(src.merged_cells.ranges):
                try:
                    if rng.max_col <= maxc:
                        dst.merge_cells(str(rng))
                except Exception:
                    pass
            for letter, dim in src.column_dimensions.items():
                d2 = dst.column_dimensions[letter]
                d2.width = dim.width
                d2.hidden = dim.hidden
            for rn, dim in src.row_dimensions.items():
                if dim.height:
                    dst.row_dimensions[rn].height = dim.height
            dst.row_dimensions[1].height = src.row_dimensions[1].height if src.row_dimensions[1].height else dst.row_dimensions[1].height
            # view / color / filter / freeze
            dst.sheet_view.rightToLeft = True
            dst.sheet_view.showGridLines = False
            dst.sheet_state = "visible" if (is_input or s in ALWAYS_VISIBLE or s == "داشبورد") else "hidden"
            if src.sheet_properties.tabColor:
                dst.sheet_properties.tabColor = src.sheet_properties.tabColor
            if src.auto_filter.ref:
                try:
                    a1, a2 = src.auto_filter.ref.split(":")
                    r2 = int(re.sub(r"\D", "", a2)); 
                    if r2 <= maxr:
                        dst.auto_filter.ref = src.auto_filter.ref
                except Exception:
                    pass
            if src.freeze_panes:
                dst.freeze_panes = src.freeze_panes
            # conditional formatting (rules are self-contained)
            for rng_str, rules in list(src.conditional_formatting._cf_rules.items()):
                for rule in rules:
                    try:
                        dst.conditional_formatting.add(str(rng_str), deepcopy(rule))
                    except Exception:
                        pass
            # data validations (input sheets only — snapshot cells all locked)
            if is_input:
                for dv in src.data_validations.dataValidation:
                    nd = DataValidation(type=dv.type, operator=dv.operator,
                                        formula1=dv.formula1, formula2=dv.formula2,
                                        allow_blank=True, showErrorMessage=dv.showErrorMessage,
                                        errorTitle=dv.errorTitle, error=dv.error,
                                        showInputMessage=dv.showInputMessage,
                                        promptTitle=dv.promptTitle, prompt=dv.prompt)
                    dst.add_data_validation(nd)
                    for sq in dv.sqref.ranges:
                        nd.add(str(sq))
            # protection: whole file locked with dept token; inputs keep unlocked cells
            p = dst.protection
            p.sheet = True
            p.formatCells = False; p.formatColumns = False; p.formatRows = False
            p.insertColumns = False; p.insertRows = False
            p.deleteColumns = False; p.deleteRows = False
            p.sort = False; p.autoFilter = False
            p.selectLockedCells = False; p.selectUnlockedCells = False
            p.objects = False; p.scenarios = True
            p.password = TOKENS[dk]
        # names needed by formulas/DV in the dept file
        CAL_F, CAL_L = meta["cal"]["first_row"], meta["cal"]["last_row"]
        names = {
            "CalS": f"'تقویم'!$A${CAL_F}:$A${CAL_L}",
            "CalJ": f"'تقویم'!$B${CAL_F}:$B${CAL_L}",
            "Hols": f"'تنظیمات'!$D${meta['settings']['hol_rows'][0]}:$D${meta['settings']['hol_rows'][1]}",
            "HolsNorm": f"'تنظیمات'!$C${meta['settings']['hol_rows'][0]}:$C${meta['settings']['hol_rows'][1]}",
            "ParamThresh": "'تنظیمات'!$B$8", "FX_USD": "'تنظیمات'!$B$6", "FX_EUR": "'تنظیمات'!$B$7",
            "BomCodes": "'مهندسی_BOM'!$B$8:$B$31",
        }
        for nm, (rr, c1, c2) in meta["list_ranges"].items():
            if nm == "ListUsers":
                continue
            from openpyxl.utils import get_column_letter as gl
            names[nm] = f"'تنظیمات'!${gl(c1)}${rr}:${gl(c2)}${rr}"
        for nm, ref in names.items():
            dwb.defined_names.add(DefinedName(nm, attr_text=ref))
        # structure lock with dept token (dept cannot unhide snapshots)
        dwb.security = WorkbookProtection(lockStructure=True, workbookPassword=TOKENS[dk])
        dwb.active = 0
        out_dir = os.path.join(args.out, folder)
        os.makedirs(out_dir, exist_ok=True)
        out_path = os.path.join(out_dir, dfile)
        dwb.save(out_path)
        with open(os.path.join(out_dir, "README.txt"), "w", encoding="utf-8") as fh:
            fh.write(
                f"برگه‌های فعال این فایل: {', '.join(input_sheets)}\n"
                "سلول‌های زرد = ورودی شما؛ سایر سلول‌ها قفل‌اند. پس از ذخیره، مدیر پروژه با ماکروی Borna_Collect\n"
                "داده شما را در فایل مادر ادغام و با Borna_Publish اسنپ‌شات‌های بالادستی را تازه می‌کند.\n"
                "این فایل را در همین پوشه نگه دارید (ACL شبکه، نوشتن را فقط برای گروه همین واحد و PMO باز می‌گذارد).\n"
                "نیازی به فعال‌سازی ماکرو ندارد — نسخه واحد بدون ماکرو است.\n")
        summary.append((dk, dfile, input_sheets, len(snap_sheets)))
        print(f"  {dk:7} → {out_path}  (inputs: {', '.join(input_sheets)}; snapshots hidden: {len(snap_sheets)})")

    # master-side helper: ACL hints with per-dept tokens
    with open(os.path.join(args.out, "_ACL_HINTS.txt"), "w", encoding="utf-8") as fh:
        fh.write("رمز قفل فایل هر واحد (توکن) — فقط در اختیار IT و PMO:\n")
        for dk, tok in TOKENS.items():
            fh.write(f"  {dk:7} {tok}\n")
        fh.write("ساختار: icacls برای گروه Borna-<DEPT> = Modify روی پوشه خودش و Read روی پوشه‌های دیگر ممنوع؛ "
                 "Borna-PMO = Modify روی همه + master؛ دستور آماده: tools/deploy_acl.ps1\n")
    print("done:", len(summary), "unit files →", args.out)


if __name__ == "__main__":
    main()
