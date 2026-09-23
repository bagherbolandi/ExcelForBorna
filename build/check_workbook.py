#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Structural checks for the generated Borna PM workbook (no Excel needed)."""
import re
import sys
import zipfile
from openpyxl import load_workbook

PATH = sys.argv[1] if len(sys.argv) > 1 else "out/Project_Management_Borna.xlsx"
issues = []

# 1) zip integrity
zf = zipfile.ZipFile(PATH)
bad = zf.testzip()
if bad:
    issues.append(f"zip corrupt: {bad}")
for n in zf.namelist():
    zf.read(n)  # decode crc

wb = load_workbook(PATH)
sheet_names = set(wb.sheetnames)
print("sheets:", wb.sheetnames)

# 2) defined names resolve
names = {k: wb.defined_names[k].attr_text for k in wb.defined_names}
for nm, ref in names.items():
    m = re.match(r"^'([^']+)'!", ref)
    if m and m.group(1) not in sheet_names:
        issues.append(f"defined name {nm} -> missing sheet {m.group(1)}")
print(f"defined names: {len(names)}")

# 3) formula sanity across all sheets
f_re = re.compile(r"'([^']+)'!")
fnames = set()
nf = 0
for ws in wb.worksheets:
    for row in ws.iter_rows():
        for cell in row:
            v = cell.value
            if isinstance(v, str) and v.startswith("="):
                nf += 1
                if v.count("(") != v.count(")"):
                    issues.append(f"{ws.title}!{cell.coordinate}: unbalanced parens: {v[:90]}")
                if v.count('"') % 2:
                    issues.append(f"{ws.title}!{cell.coordinate}: odd quotes: {v[:90]}")
                for sh in f_re.findall(v):
                    if sh not in sheet_names:
                        issues.append(f"{ws.title}!{cell.coordinate}: bad sheet ref {sh!r}")
                for fn in re.findall(r"([A-Za-z_\.]+)\(", v):
                    fnames.add(fn)
            elif isinstance(v, str) and re.match(r"^(IF|INDEX|MATCH|SUM|ROUND|SUBSTITUTE)\(", v):
                issues.append(f"{ws.title}!{cell.coordinate}: formula lost '=' -> {v[:60]}")
print("formula cells:", nf)
unexpected = fnames - {"IF", "IFERROR", "INDEX", "MATCH", "SUM", "SUMPRODUCT", "SUMIF", "SUMIFS",
                       "COUNT", "COUNTA", "COUNTIF", "COUNTIFS", "ROUND", "N", "AND", "OR", "NOT",
                       "ISNUMBER", "SEARCH", "TRIM", "TEXT", "SUBSTITUTE", "TODAY", "WEEKDAY",
                       "NETWORKDAYS", "MEDIAN", "MAX", "MIN", "NETWORKDAYS.INTL", "INT"}
if unexpected:
    print("WARNING: unexpected functions:", unexpected)

# 4) per-sheet structural expectations
expect = {
    "گانت": {"cf_min": 6, "tables": True},
    "داشبورد": {"charts": 4},
    "مهندسی_BOM": {"rows_min": 32},
}
ws = wb["گانت"]
print("gantt CF ranges:", len(ws.conditional_formatting._cf_rules))
for r in ws.conditional_formatting._cf_rules.values():
    pass
tot_cf = sum(len(v) for v in ws.conditional_formatting._cf_rules.values())
if tot_cf < 6:
    issues.append(f"gantt CF rules only {tot_cf}")
ws = wb["داشبورد"]
print("dashboard charts:", len(ws._charts))
if len(ws._charts) < 4:
    issues.append(f"dashboard charts {len(ws._charts)} < 4")
# grid header: serial row present?
if wb["گانت"]["T6"].value is None:
    issues.append("gantt serial header row T6 empty")
# 5) protection state
locked_all = all(ws.protection.sheet for ws in wb.worksheets)
if not locked_all:
    issues.append("some sheets unprotected")
print("all sheets protected:", locked_all, "| structure locked:",
      wb.security.lockStructure if wb.security else False)
# inputs unlocked sample: شناسنامه B5 must be unlocked
for sh, coord in [("شناسنامه", "B5"), ("گانت", "F9"), ("بازرگانی", "I8"), ("تنظیمات", "A53"), ("تصویب", "B12")]:
    c = wb[sh][coord]
    if c.protection.locked:
        issues.append(f"{sh}!{coord} should be unlocked (input)")
# formula cells must be locked
for sh, coord in [("گانت", "K8"), ("مهندسی_BOM", "H8"), ("بازرگانی", "N8"), ("داشبورد", "A6")]:
    c = wb[sh][coord]
    if not c.protection.locked:
        issues.append(f"{sh}!{coord} formula cell should be locked")

# 6) RTL + font spot checks
if not wb["گانت"].sheet_view.rightToLeft:
    issues.append("gantt not RTL")
f = wb["بازرگانی"]["C8"].font
print("sample font:", f.name)
if f.name != "Tahoma":
    issues.append("non-Tahoma font found")

# 7) DV coverage
for sh in ["گانت", "ریسک", "فروش", "ورود"]:
    dv = wb[sh].data_validations.dataValidation
    print(f"DV {sh}:", len(dv))

# 8) calendar row count + lookup key format
c = wb["تقویم"]["B8"]
print("cal first:", wb["تقویم"]["A8"].value, c.value)
last = 0
for row in wb["تقویم"].iter_rows(min_row=8, min_col=2, max_col=2):
    if row[0].value:
        last = row[0].row
if last < 2000:
    issues.append("calendar rows short")
print("cal last row:", last)

print("\n" + ("ISSUES FOUND:\n" if issues else "ALL CHECKS PASSED ✅"))
for i in issues:
    print(" -", i)
sys.exit(1 if issues else 0)
