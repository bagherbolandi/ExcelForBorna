"""
model.py — in-memory workbook model
====================================
A pure-Python description of the workbook. The builder writes this model:
  - to a real .xlsx via openpyxl (renderer)
  - to the formula engine for verification (same formulas, same values)

A single formula string is stored per cell; it is valid A1-syntax Excel that
the engine also understands natively. (Structured-reference showcase formulas
on Dashboard/_Trace are translated to A1 by a small translator in verify.py,
so file & engine stay in lockstep.)
"""
from __future__ import annotations


class CellVal:
    __slots__ = ("value", "formula")

    def __init__(self, value=None, formula=None):
        self.value = value
        self.formula = formula

    @property
    def is_formula(self) -> bool:
        return bool(self.formula)


class TableDef:
    """Description of an Excel table placed on a sheet."""
    def __init__(self, name, sheet, header_row, first_data_row):
        self.name = name
        self.sheet = sheet
        self.header_row = header_row          # 1-based
        self.first_data_row = first_data_row  # 1-based
        self.last_data_row = first_data_row   # updated as rows are added
        self.n_cols = 0


class SheetModel:
    def __init__(self, name: str):
        self.name = name
        self.rows: dict[int, dict[int, CellVal]] = {}   # 1-based both
        self.tables: list[TableDef] = []
        self.merged: list[tuple[int, int, int, int]] = []  # r1,c1,r2,c2 (1-based)
        self.col_widths: dict[int, float] = {}
        self.freeze = "A2"

    # ---- cell helpers ----
    def set(self, r, c, value=None, formula=None):
        self.rows.setdefault(r, {})[c] = CellVal(value, formula)

    def get(self, r, c) -> CellVal | None:
        return self.rows.get(r, {}).get(c)

    def cell_ref(self, r, c):
        from openpyxl.utils import get_column_letter
        return f"{get_column_letter(c)}{r}"

    def max_col(self) -> int:
        m = 0
        for row in self.rows.values():
            for c in row:
                m = max(m, c)
        return m

    def max_row(self) -> int:
        return max(self.rows.keys()) if self.rows else 0

    def add_table(self, name, header_row, first_data_row):
        t = TableDef(name, self.name, header_row, first_data_row)
        self.tables.append(t)
        return t


class WorkbookModel:
    def __init__(self):
        self.sheets: dict[str, SheetModel] = {}
        self.order: list[str] = []

    def sheet(self, name: str) -> SheetModel:
        if name not in self.sheets:
            self.sheets[name] = SheetModel(name)
            self.order.append(name)
        return self.sheets[name]

    def has(self, name: str) -> bool:
        return name in self.sheets


# --------------------------------------------------------------------------- #
# Column-letter helper (independent of openpyxl for the engine side)
# --------------------------------------------------------------------------- #
_COLS = "ABCDEFGHIJKLMNOPQRSTUVWXYZ"


def col_letter(idx: int) -> str:
    """1-based column index -> letters (A, B, ... Z, AA ...)."""
    s = ""
    while idx > 0:
        idx, rem = divmod(idx - 1, 26)
        s = _COLS[rem] + s
    return s


def ref_ok(r, c):
    return f"{col_letter(c)}{r}"
